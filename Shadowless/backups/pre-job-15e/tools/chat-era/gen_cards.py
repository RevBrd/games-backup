#!/usr/bin/env python3
"""Generate cards.js (card DB + deck lists) straight from source JSON. No hand-typing."""
import json, glob, collections
from openpyxl import load_workbook

TYPE = {'Grass':'G','Fire':'R','Water':'W','Lightning':'L','Psychic':'P',
        'Fighting':'F','Darkness':'D','Metal':'M','Colorless':'C'}

CARDS = {}
for f in glob.glob('/tmp/ptcg/*.json'):
    for c in json.load(open(f)):
        CARDS[c['id']] = c

wb = load_workbook('/mnt/user-data/uploads/P_TCG_Data.xlsx', read_only=True)
WANT = ['Brushfire', 'Overgrowth', 'Blackout', 'Zap']
EXTRA_IDS = ['base1-3', 'base1-9', 'base1-11', 'base1-13', 'base1-14', 'base1-16', 'base1-18', 'base1-19', 'base1-20', 'base1-25', 'base1-26', 'base1-36', 'base1-37', 'base1-40', 'base1-41', 'base1-45', 'base1-47', 'base1-48', 'base1-51', 'base1-54', 'base1-59', 'base1-60', 'base1-61', 'base1-67', 'base1-72', 'base1-73', 'base1-74', 'base1-76', 'base1-77', 'base1-78', 'base1-82', 'base1-83', 'base1-85', 'base1-86', 'base1-87', 'base1-89', 'base1-96']

decks = {}
order = []
for n in WANT:
    rows = []
    for r in wb[n].iter_rows(values_only=True):
        if not r or r[0] is None or str(r[0]).strip().lower() == 'quantity':
            continue
        try: q = int(float(r[0]))
        except (TypeError, ValueError): continue
        cid = str(r[2]).strip()
        rows.append([q, cid])
        if cid not in order: order.append(cid)
    decks[n] = rows

# Cards implemented but not (yet) in any deck still belong in the database, so
# the deck generator and collection can use them.
for cid in EXTRA_IDS:
    if cid not in order: order.append(cid)

def js(s):
    return json.dumps(s, ensure_ascii=False)

out = []
out.append('// AUTO-GENERATED from pokemon-tcg-data + P_TCG_Data.xlsx. Do not hand-edit.')
out.append('const CARD_DB = {')
for cid in order:
    c = CARDS[cid]
    st = c['supertype']
    o = {'id': cid, 'name': c['name'], 'set': cid.split('-')[0], 'num': c['number'],
         'rarity': c.get('rarity', ''), 'artist': c.get('artist', '')}
    if st == 'Pok\u00e9mon':
        stage = next((x for x in ['Baby','Stage 2','Stage 1','Basic'] if x in c.get('subtypes',[])), 'Basic')
        w = (c.get('weaknesses') or [{}])[0]
        r = (c.get('resistances') or [{}])[0]
        o.update(kind='pokemon', stage=stage, hp=int(c['hp']),
                 type=TYPE.get((c.get('types') or ['Colorless'])[0], 'C'),
                 evolvesFrom=c.get('evolvesFrom', ''),
                 wkType=TYPE.get(w.get('type',''), ''), wkVal=(w.get('value') or '').replace('\u00d7','x'),
                 rsType=TYPE.get(r.get('type',''), ''), rsVal=(r.get('value') or ''),
                 retreat=len(c.get('retreatCost') or []))
        o['attacks'] = [{'name': a['name'],
                         'cost': ''.join(TYPE.get(x,'C') for x in (a.get('cost') or [])),
                         'dmg': a.get('damage',''),
                         'text': a.get('text','')} for a in c.get('attacks', [])]
        ab = (c.get('abilities') or [])
        o['power'] = {'name': ab[0]['name'], 'text': ab[0]['text'], 'type': ab[0]['type']} if ab else None
    elif st == 'Trainer':
        sub = next((x for x in c.get('subtypes',[]) if x in ('Stadium','Pok\u00e9mon Tool')), 'Trainer')
        o.update(kind='trainer', sub=sub, text=' '.join(c.get('rules', [])))
    else:
        cls = 'Special' if 'Special' in c.get('subtypes',[]) else 'Basic'
        # Basic Energy cards carry NO "types" field in the source data - the
        # provided type has to come from the card name.
        provides = ''.join(TYPE.get(x,'C') for x in (c.get('types') or []))
        if not provides:
            base = c['name'].replace(' Energy','').strip()
            provides = TYPE.get(base, '')
        if not provides:
            # Special Energy that provides more than one symbol
            SPECIAL = {'Double Colorless Energy': 'CC'}
            provides = SPECIAL.get(c['name'], '')
        if not provides:
            raise SystemExit(f'UNRESOLVED energy type for {cid} "{c["name"]}"')
        o.update(kind='energy', cls=cls, provides=provides,
                 text=' '.join(c.get('rules', [])))
    out.append('  ' + js(cid) + ': ' + js(o) + ',')
out.append('};')
out.append('')
out.append('const DECKS = {')
for n, rows in decks.items():
    out.append(f'  {js(n)}: {{ name: {js(n)}, list: [')
    for q, cid in rows:
        out.append(f'    [{q}, {js(cid)}],  // {CARDS[cid]["name"]}')
    out.append('  ]},')
out.append('};')
out.append('')
out.append("if (typeof module !== 'undefined') module.exports = { CARD_DB, DECKS };")
out.append('')

src = '\n'.join(out) + '\n'
open('/tmp/ptcg/build/cards.js', 'w', encoding='utf-8').write(src)
print('wrote cards.js:', len(src), 'bytes,', len(order), 'cards,', len(decks), 'decks')
for n, rows in decks.items():
    print(' ', n, sum(q for q, _ in rows), 'cards')
